"""Build the numbered clip catalog: assets/packs/catalog.json + CLIP_CATALOG.docx (with pictures).

usage (from pov-meme-kit/): python3 scripts/make_catalog.py
Clip ids: C1..C29 (cat pack), D1..D29 (dog pack). Stories use them as actor "clip": "C7".
"""
import json, os, subprocess, sys, tempfile
import numpy as np
from PIL import Image, ImageOps
from docx import Document
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.dirname(__file__))
from catalog_segments import CAT_CUTS, DOG_CUTS, segments

CAT = {  # id: (description, mood / good for)
    1: ("White cat with huge round 'huh?' eyes peeking up", "shock, confusion, 'huh?'"),
    2: ("Fluffy orange cat sitting, looking up", "hoping, waiting, wishing"),
    3: ("Small orange cat sitting and meowing", "calling out, asking"),
    4: ("White & brown cat standing, meowing loudly", "demanding, complaining, kid asking"),
    5: ("Same white & brown cat, close, meowing", "pleading, arguing"),
    6: ("Big-eyed orange kitten peeking from the side", "innocent child, curious, worried (audio: a woman laughing then talking - user flagged)"),
    7: ("Sad tabby (calico) looking down", "sad, tired, guilty, quiet sacrifice"),
    8: ("Cream cat standing on hind legs, paws together", "polite, begging, cute request"),
    9: ("Angry white cat and tabby hissing/screaming (rapid cuts)", "angry, yelling, fighting"),
    10: ("Black & white kitten wrapped in a towel, held", "baby, newborn, fragile"),
    11: ("Orange cat turning away, being held", "ignoring, upset, walking away"),
    12: ("Small dark cat jumping/climbing far away", "tiny in distance, playing"),
    13: ("Cat peeking from the edge with sad big eyes", "spying, secretly watching, worried"),
    14: ("Grey striped kitten with paws up, screaming", "panicking, tantrum, crying out"),
    15: ("Chubby white cat sitting, small in frame", "waiting, lonely, calm"),
    16: ("Dark tabby standing still, staring", "serious, stern, shocked silence"),
    17: ("Scottish fold behind a cream cat", "comforting, hugging, together"),
    18: ("Dark cat crying, tongue out, blurry close-up", "crying, heartbroken, tears of joy"),
    19: ("White cat standing and waving/dancing", "happy, excited, celebrating, greeting"),
    20: ("Cat sitting eating from a metal bowl", "eating, hungry, leftovers"),
    21: ("White cat driving a car (steering wheel) + pop cat (mixed)", "driving, picking someone up"),
    22: ("Orange cat close-up, sleepy/annoyed face", "tired, unimpressed"),
    23: ("Orange cat 'singing', eyes half-closed", "talking, singing, proud (audio is a human singing)"),
    24: ("Cat in a blue cap standing on hind legs", "worker, delivery, playful"),
    25: ("Tabby in a costume dancing (pink wings/skirt)", "party, celebration, silly"),
    26: ("Tabby sitting up with a shocked face", "shocked, surprised, frozen"),
    27: ("Teary-eyed tabby held up by hands", "heartbroken, sad, about to cry"),
    28: ("Fluffy white cat lying, rolling", "relaxed, lazy, comfy"),
    29: ("Brown cat standing on hind legs", "waiting, standing guard, determined"),
}
DOG = {
    1: ("Fluffy white chihuahua, eyes closed, standing", "smug, proud, ignoring"),
    2: ("Jack Russell jumping with tongue out", "happy, excited, running to greet"),
    3: ("German shepherd with a butterfly on its nose", "calm, gentle, butterfly / spirit"),
    4: ("Husky nose close-up staring into camera", "curious, sniffing, shocked zoom"),
    5: ("Small dog wrapped in a blanket", "sick, cold, cozy, cared for"),
    6: ("Chihuahua squinting, tongue out", "disgusted, sassy, teasing"),
    7: ("Black dog side-eye, close", "suspicious, judging"),
    8: ("Pug walking", "walking away, arriving"),
    9: ("Rottweiler-type dog with a bow tie, sitting", "groom, formal, polite gentleman"),
    10: ("Black & white dog sitting, serious", "serious, listening"),
    11: ("Bug-eyed chihuahua close-up", "shocked, scared, 'what?!'"),
    12: ("Small white dog grinning with teeth", "nervous smile, awkward"),
    13: ("Corgi smiling, small in frame", "happy, friendly"),
    14: ("Chocolate lab with a goofy face", "funny, silly, surprised"),
    15: ("Corgi puppy eating watermelon", "eating, hungry, happy child"),
    16: ("Beagle grimace/smile close-up", "awkward, guilty grin"),
    17: ("Dog in a straw hat talking on a phone", "phone call, farmer, grandpa"),
    18: ("Tan dog close-up, serious stare", "stern, disappointed"),
    19: ("Pointer standing on long legs", "standing, waiting, guard"),
    20: ("Same pointer, closer", "standing, waiting"),
    21: ("Husky howling upward", "crying out, singing, missing someone"),
    22: ("Boston terrier sitting with collar", "polite, waiting, listening"),
    23: ("Dog standing on hind legs (from behind)", "begging, reaching up"),
    24: ("Tiny fluffy puppy far away", "baby, small, lost"),
    25: ("Small black dog walking", "small, lonely, walking"),
    26: ("Bull terrier in a green hood", "cozy, silly, child"),
    27: ("Dog sitting like a person, leaning", "relaxed, boss, chilling"),
    28: ("Golden retriever with headset and tie", "office worker, call center, doctor, boss"),
    29: ("White dog sitting and howling", "crying, singing, lonely"),
}


def sound_type(src, a, b):
    raw = subprocess.run(["ffmpeg", "-v", "error", "-ss", str(a), "-to", str(b), "-i", src, "-vn", "-ac", "1",
                          "-ar", "16000", "-f", "f32le", "-"], capture_output=True).stdout
    x = np.frombuffer(raw, np.float32)
    if len(x) == 0 or np.sqrt((x ** 2).mean()) < 0.005:
        return "silent"
    f0s, frames = [], 0
    for i in range(0, len(x) - 1024, 512):
        fr = x[i:i + 1024] * np.hanning(1024)
        frames += 1
        if np.sqrt((fr ** 2).mean()) < 0.01:
            continue
        ac = np.correlate(fr, fr, "full")[1023:]
        lag = 13 + np.argmax(ac[13:200])
        if ac[lag] / ac[0] > 0.5:
            f0s.append(16000 / lag)
    if len(f0s) < 5:
        return "noise / effects (no clear voice)"
    m = np.median(f0s)
    return f"animal sound (~{m:.0f} Hz)" if m >= 380 else f"low voice ~{m:.0f} Hz (human or bark - check)"


def main():
    packs = {"C": ("cat", "assets/packs/cat_pack.mp4", CAT_CUTS, CAT), "D": ("dog", "assets/packs/dog_pack.mp4", DOG_CUTS, DOG)}
    tmp = tempfile.mkdtemp()
    catalog, doc = {}, Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Meme Clip Catalog — cats (C) and dogs (D)", 0)
    doc.add_paragraph("Every usable green-screen clip, numbered. When planning a video, Claude lists the clip "
                      "numbers it will use (e.g. 'Mom = C7, Me = C6, Doctor = D28'). Pictures are the middle "
                      "frame of each clip. 'Sound' is what the clip's own audio is — the video plays it while "
                      "that character talks.")
    for pre, (name, src, cuts, desc) in packs.items():
        doc.add_heading(f"{name.capitalize()} clips", 1)
        for n, (a, b) in enumerate(segments(cuts), 1):
            cid = f"{pre}{n}"
            d, mood = desc[n]
            snd = sound_type(src, a + 0.5, b - 0.2)
            catalog[cid] = {"pack": name, "start": a, "end": b, "description": d, "good_for": mood, "sound": snd}
            png, img = f"{tmp}/{cid}.png", f"{tmp}/{cid}.jpg"
            subprocess.run(["ffmpeg", "-v", "error", "-y", "-ss", str((a + b) / 2), "-i", src, "-frames:v", "1",
                            "-vf", "format=rgba,chromakey=0x50BB4A:0.10:0.04,despill=type=green", png], check=True)
            cut = Image.open(png).convert("RGBA")
            flat = Image.new("RGB", cut.size, "white")
            flat.paste(cut, mask=cut.split()[3])  # cut-out on white, no transparency (shows in every viewer)
            flat.crop(flat.getbbox() and ImageOps.invert(flat).getbbox() or (0, 0, *flat.size)).save(img, quality=88)
            doc.add_heading(f"{n}. Clip {cid} — {d}", 2)
            doc.add_picture(img, width=Inches(3.2))
            doc.add_paragraph(f"Good for: {mood}\nTime in {name}_pack.mp4: {a:.2f}–{b:.2f} s ({b - a:.1f} s)\nSound: {snd}")
            print(cid, snd, flush=True)
    json.dump(catalog, open("assets/packs/catalog.json", "w"), indent=1)
    doc.save("CLIP_CATALOG.docx")


main()
